import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE = "medical_billing.xlsx"

# Create Excel file
if os.path.exists(FILE):
    wb = load_workbook(FILE)
    ws = wb.active
else:
    wb = Workbook()
    ws = wb.active

if ws.max_row == 1 and ws["A1"].value is None:
    ws.append(["ID","Patient","Doctor","Med",
               "Qty","Price","Total","Pay"])

wb.save(FILE)


def dashboard():
    win = tk.Tk()
    win.title("Medical Billing")
    win.geometry("360x600")

    bill = tk.StringVar()
    patient = tk.StringVar()
    doctor = tk.StringVar()
    medicine = tk.StringVar()
    qty = tk.StringVar()
    price = tk.StringVar()
    total = tk.StringVar()
    payment = tk.StringVar(value="Cash")

    def clear():
        for x in [bill, patient, doctor, medicine,
                  qty, price, total]:
            x.set("")

    def calculate():
        try:
            total.set(int(qty.get()) * float(price.get()))
            return True
        except:
            messagebox.showerror("Error", "Enter valid values")
            return False

    def add():
        if not calculate():
            return

        wb = load_workbook(FILE)
        ws = wb.active

        ws.append([
            bill.get(), patient.get(), doctor.get(),
            medicine.get(), qty.get(), price.get(),
            total.get(), payment.get()
        ])

        wb.save(FILE)

        messagebox.showinfo("Success", "Bill Added")
        view()
        clear()

    def view():
        table.delete(*table.get_children())

        wb = load_workbook(FILE)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            table.insert("", "end", values=row)

    tk.Label(
        win,
        text="MEDICAL BILLING",
        font=("Arial", 13, "bold")
    ).pack(pady=5)

    form = tk.Frame(win)
    form.pack()

    fields = [
        ("ID", bill),
        ("Patient", patient),
        ("Doctor", doctor),
        ("Med", medicine),
        ("Qty", qty),
        ("₹", price),
        ("Total", total)
    ]

    for i, (name, var) in enumerate(fields):
        tk.Label(form, text=name).grid(row=i, column=0)
        tk.Entry(
            form,
            textvariable=var,
            width=16
        ).grid(row=i, column=1, pady=1)

    tk.Label(form, text="Pay").grid(row=7, column=0)

    ttk.Combobox(
        form,
        textvariable=payment,
        values=["Cash", "UPI", "Card"],
        width=13
    ).grid(row=7, column=1)

    tk.Button(
        win,
        text="Calculate / Add",
        command=add,
        width=15
    ).pack(pady=3)

    tk.Button(
        win,
        text="Clear",
        command=clear,
        width=15
    ).pack(pady=2)

    table = ttk.Treeview(
        win,
        columns=(
            "ID", "Patient", "Doctor", "Med",
            "Qty", "Price", "Total", "Pay"
        ),
        show="headings",
        height=4
    )

    for c in table["columns"]:
        table.heading(c, text=c)
        table.column(c, width=75)

    table.pack(fill="x", padx=3, pady=5)

    view()
    win.mainloop()


def login():
    if user.get() == "admin" and password.get() == "1234":
        login_win.destroy()
        dashboard()
    else:
        messagebox.showerror("Error", "Invalid Login")


login_win = tk.Tk()
login_win.title("Login")
login_win.geometry("280x220")

tk.Label(
    login_win,
    text="MEDICAL BILLING",
    font=("Arial", 14, "bold")
).pack(pady=15)

tk.Label(login_win, text="User").pack()

user = tk.Entry(login_win)
user.pack()

tk.Label(login_win, text="Pass").pack()

password = tk.Entry(login_win, show="*")
password.pack()

tk.Button(
    login_win,
    text="LOGIN",
    command=login,
    width=10
).pack(pady=15)

login_win.mainloop()
